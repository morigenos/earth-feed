"""EU fuel price collector.

fixtures/wob_history_sample.xlsx is the Commission's own history workbook cut down to
six bulletin dates (5 Oct 2026, 28 Sep, 21 Sep, 21 Dec 2020, 14 Dec 2020, 3 Jan 2005),
both price sheets, original headers. Source: European Commission, Weekly Oil Bulletin;
reproduction authorised provided the source is acknowledged.
"""
import io, sys, json, pathlib, tempfile, unittest, datetime
from unittest.mock import patch, Mock
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / 'scripts'))
import feed_health as health
import fetch_fuel_eu as fuel

FIXTURE = pathlib.Path(__file__).resolve().parent / 'fixtures' / 'wob_history_sample.xlsx'


def edited(change):
    """Fixture bytes after applying change(workbook)."""
    import openpyxl
    wb = openpyxl.load_workbook(FIXTURE)
    change(wb)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def column(ws, name):
    for cell in ws[1]:
        if cell.value == name: return cell.column
    raise KeyError(name)


class FuelTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old = health.OUT
        health.OUT = pathlib.Path(self.tmp.name)
        self.blob = FIXTURE.read_bytes()

    def tearDown(self):
        health.OUT = self.old
        self.tmp.cleanup()

    def data(self, name): return json.loads((health.OUT / (name + '.json')).read_text())

    def item(self, iso3, fuel_type):
        return next(i for i in self.data('fuel')['items'] if i[0] == iso3 and i[1] == fuel_type)

    def test_real_layout_parses_and_publishes(self):
        self.assertTrue(health.run('fuel', lambda: fuel.main(self.blob)))
        j = self.data('fuel')
        self.assertEqual(j['validTime'], '2026-10-05')
        self.assertEqual(j['class'], 'reported')
        self.assertEqual(len({i[0] for i in j['items']}), 27)
        self.assertEqual(self.item('SVN', 'GASOLINE_95')[2:5], [1.763, 1.005, '2026-10-05'])
        self.assertEqual(self.data('health')['datasets']['fuel']['status'], 'ok')

    def test_greece_and_uk_codes_map_to_iso3(self):
        fuel.main(self.blob)
        areas = {i[0] for i in self.data('fuel')['items']}
        self.assertIn('GRC', areas)
        self.assertNotIn('GR', areas)
        self.assertTrue((health.OUT / 'fuel' / 'history' / 'GBR.json').exists())

    def test_uk_is_history_only_after_2020(self):
        fuel.main(self.blob)
        j = self.data('fuel')
        self.assertNotIn('GBR', {i[0] for i in j['items']})
        self.assertIn(['GBR', '2020-12-21'], j['stale'])

    def test_local_currency_for_non_euro_members(self):
        fuel.main(self.blob)
        hu = self.item('HUN', 'GASOLINE_95')
        self.assertEqual(hu[7], 'HUF')
        self.assertAlmostEqual(hu[8], 1725.568103447765 / 0.002718868950516585 / 1000, places=2)
        self.assertEqual(self.item('SVN', 'DIESEL')[7], 'EUR')

    def test_history_is_ascending_with_nulls_not_zeros(self):
        fuel.main(self.blob)
        h = json.loads((health.OUT / 'fuel' / 'history' / 'HRV.json').read_text())
        self.assertEqual(h['dates'], sorted(h['dates']))
        self.assertEqual(h['dates'][0], 20050103)
        self.assertIsNone(h['series']['GASOLINE_95']['tax'][0])   # Croatia reported from mid-2013
        self.assertEqual(h['unit'], 'EUR/L')
        self.assertIn('European Commission', h['attribution'])

    def test_aggregates_are_reference_not_countries(self):
        fuel.main(self.blob)
        j = self.data('fuel')
        self.assertNotIn('EU', {i[0] for i in j['items']})
        self.assertEqual(j['ref']['EU']['GASOLINE_95'][0], 2.006)

    def test_implausible_value_is_dropped_and_counted(self):
        def spoil(wb):
            ws = wb['Prices with taxes']
            ws.cell(row=4, column=column(ws, 'SI_price_with_tax_euro95'), value=99999)
        fuel.main(edited(spoil))
        si = self.item('SVN', 'GASOLINE_95')
        self.assertEqual(si[4], '2026-09-28')                 # falls back to the previous valid week
        self.assertIn('1 implausible', self.data('fuel')['note'])

    def test_too_few_countries_keeps_previous_file(self):
        fuel.main(self.blob)
        before = (health.OUT / 'fuel.json').read_bytes()
        def blank(wb):
            for name in ('Prices with taxes', 'Prices wo taxes'):
                ws = wb[name]
                for c in ws[1]:
                    if c.value and '_price_' in str(c.value) and not str(c.value).startswith(('AT_', 'BE_', 'EU_', 'EUR_')):
                        for r in range(4, ws.max_row + 1):
                            ws.cell(row=r, column=c.column).value = None   # cell(value=None) is a no-op
        self.assertFalse(health.run('fuel', lambda: fuel.main(edited(blank))))
        self.assertEqual(before, (health.OUT / 'fuel.json').read_bytes())
        self.assertEqual(self.data('health')['datasets']['fuel']['status'], 'failed')

    def test_sheets_must_agree_on_dates(self):
        def shift(wb): wb['Prices wo taxes'].cell(row=4, column=1, value=datetime.datetime(2026, 10, 6))
        with self.assertRaises(ValueError): fuel.parse(edited(shift))

    def test_missing_header_rejected(self):
        def wreck(wb):
            for c in wb['Prices with taxes'][1]: c.value = None
        with self.assertRaises(ValueError): fuel.parse(edited(wreck))

    def test_future_bulletin_rejected(self):
        def future(wb):
            for name in ('Prices with taxes', 'Prices wo taxes'):
                wb[name].cell(row=4, column=1, value=datetime.datetime(2099, 1, 5))
        with self.assertRaises(ValueError): fuel.parse(edited(future))

    def test_html_instead_of_workbook_is_failure(self):
        page = Mock(text='<a href="/x?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx">x</a>')
        page.raise_for_status = Mock()
        bad = Mock(content=b'<html>maintenance</html>'); bad.raise_for_status = Mock()
        with patch.object(fuel.requests, 'get', side_effect=[page, bad]):
            self.assertFalse(health.run('fuel', fuel.main))
        self.assertFalse((health.OUT / 'fuel.json').exists())

    def test_history_link_found_on_page(self):
        html = '<a href="/document/download/abc_en?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx">'
        self.assertTrue(fuel.find_history_url(html).startswith('https://energy.ec.europa.eu/document/download/abc_en'))
        self.assertEqual(fuel.find_history_url('<a href="/other.pdf">'), fuel.FALLBACK_URL)

    def test_recent_file_skips_download(self):
        fuel.main(self.blob)
        with patch.object(fuel.requests, 'get') as get:
            self.assertTrue(health.run('fuel', fuel.main))
            get.assert_not_called()

    def test_manifest_lists_fuel(self):
        fuel.main(self.blob); health.manifest()
        self.assertIn('fuel.json', self.data('index')['files'])


if __name__ == '__main__':
    unittest.main()
